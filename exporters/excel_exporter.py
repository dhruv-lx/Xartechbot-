import os
from datetime import datetime
from pathlib import Path
from PIL import Image as PILImage
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
from config import EXPORTS_DIR, PHOTOS_DIR

def create_thumbnail(image_path, size=(90, 90)):
    if not image_path or not os.path.exists(image_path):
        return None
    try:
        thumb_dir = PHOTOS_DIR / "thumbnails"
        thumb_dir.mkdir(parents=True, exist_ok=True)
        filename = Path(image_path).name
        thumb_path = thumb_dir / f"thumb_{filename}"
        
        with PILImage.open(image_path) as img:
            img.thumbnail(size, PILImage.Resampling.LANCZOS)
            # Convert RGBA or other modes to RGB for JPEG compatibility if needed
            if img.mode in ("RGBA", "P"):
                img = img.convert("RGB")
            img.save(thumb_path, "JPEG", quality=85)
        return str(thumb_path)
    except Exception as e:
        print(f"Error creating thumbnail for {image_path}: {e}")
        return None

def generate_excel_report(visits, title="Xartech Client Visits"):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"Xartech_Visits_{timestamp}.xlsx"
    filepath = EXPORTS_DIR / filename
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Client Visits"
    
    # Enable gridlines
    ws.views.sheetView[0].showGridLines = True
    
    # Header Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    # Title Row
    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = f"XARTECH - CLIENT VISITS REPORT ({datetime.now().strftime('%d %B %Y')})"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    title_cell.alignment = center_align
    ws.row_dimensions[1].height = 40
    
    # Headers
    headers = [
        "ID", "Shop Name", "Owner Name", "Phone Number", 
        "Address / Location", "Shop Photo", "Visiting Card", 
        "Remarks / Notes", "Agent Name", "Agent Username", "Visit Date & Time"
    ]
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[2].height = 28
    
    # Column Widths
    col_widths = {
        1: 8,   # ID
        2: 24,  # Shop Name
        3: 18,  # Owner Name
        4: 16,  # Phone
        5: 28,  # Address
        6: 18,  # Shop Photo
        7: 18,  # Visiting Card
        8: 26,  # Remarks
        9: 18,  # Agent Name
        10: 16, # Agent User
        11: 20  # Date & Time
    }
    
    for col_idx, width in col_widths.items():
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    current_row = 3
    for v in visits:
        ws.row_dimensions[current_row].height = 75 # Space for thumbnail image
        
        # ID
        c = ws.cell(row=current_row, column=1, value=v.get('id', ''))
        c.alignment = center_align
        c.border = thin_border
        
        # Shop Name
        c = ws.cell(row=current_row, column=2, value=v.get('shop_name', ''))
        c.alignment = left_align
        c.font = Font(bold=True)
        c.border = thin_border
        
        # Owner Name
        c = ws.cell(row=current_row, column=3, value=v.get('owner_name') or 'N/A')
        c.alignment = left_align
        c.border = thin_border
        
        # Phone
        c = ws.cell(row=current_row, column=4, value=v.get('phone_number', ''))
        c.alignment = center_align
        c.border = thin_border
        
        # Address
        addr = v.get('address') or ''
        if v.get('latitude') and v.get('longitude'):
            addr += f" (GPS: {v.get('latitude')}, {v.get('longitude')})"
        c = ws.cell(row=current_row, column=5, value=addr if addr else 'GPS Recorded')
        c.alignment = left_align
        c.border = thin_border
        
        # Shop Photo Thumbnail
        shop_thumb = create_thumbnail(v.get('shop_photo_path'))
        c_shop = ws.cell(row=current_row, column=6)
        c_shop.border = thin_border
        if shop_thumb and os.path.exists(shop_thumb):
            try:
                img = OpenpyxlImage(shop_thumb)
                img.width = 80
                img.height = 80
                col_letter = openpyxl.utils.get_column_letter(6)
                ws.add_image(img, f"{col_letter}{current_row}")
            except Exception as e:
                c_shop.value = "Image Attached"
                c_shop.alignment = center_align
        else:
            c_shop.value = "No Photo"
            c_shop.alignment = center_align
            
        # Card Photo Thumbnail
        card_thumb = create_thumbnail(v.get('card_photo_path'))
        c_card = ws.cell(row=current_row, column=7)
        c_card.border = thin_border
        if card_thumb and os.path.exists(card_thumb):
            try:
                img = OpenpyxlImage(card_thumb)
                img.width = 80
                img.height = 80
                col_letter = openpyxl.utils.get_column_letter(7)
                ws.add_image(img, f"{col_letter}{current_row}")
            except Exception as e:
                c_card.value = "Card Attached"
                c_card.alignment = center_align
        else:
            c_card.value = "No Card"
            c_card.alignment = center_align
            
        # Remarks
        c = ws.cell(row=current_row, column=8, value=v.get('remarks') or '-')
        c.alignment = left_align
        c.border = thin_border
        
        # Agent Name
        c = ws.cell(row=current_row, column=9, value=v.get('agent_name', ''))
        c.alignment = left_align
        c.border = thin_border
        
        # Agent Username
        c = ws.cell(row=current_row, column=10, value=v.get('agent_username', ''))
        c.alignment = center_align
        c.border = thin_border
        
        # Date & Time
        c = ws.cell(row=current_row, column=11, value=str(v.get('created_at', '')))
        c.alignment = center_align
        c.border = thin_border
        
        current_row += 1

    wb.save(filepath)
    return filepath
